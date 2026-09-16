"""Unit tests for the Excel source adapter (in-memory xlsx, no network)."""

from __future__ import annotations

import io

import openpyxl
import pytest
from app.infrastructure.sources import read_source
from app.infrastructure.sources.excel import read_xlsx

pytestmark = pytest.mark.unit


def _xlsx_bytes(rows: list[list[object]], sheet: str = "Sheet1") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_simple_header_and_rows() -> None:
    data = _xlsx_bytes(
        [
            ["Name", "Price", "SKU"],
            ["Product A", 1000, "A-1"],
            ["Product B", 2000.5, "B-2"],
        ]
    )
    read = read_xlsx(data)
    assert read.complete is True
    assert read.error is None
    assert read.headers == ["Name", "Price", "SKU"]
    assert read.row_count == 2
    assert read.rows[0] == {"Name": "Product A", "Price": "1000", "SKU": "A-1"}
    assert read.rows[1] == {"Name": "Product B", "Price": "2000.5", "SKU": "B-2"}


def test_empty_rows_skipped_and_cells_stringified() -> None:
    data = _xlsx_bytes(
        [
            ["Name", "Price"],
            ["A", 1],
            [None, None],  # empty row
            ["B", 2],
        ]
    )
    read = read_xlsx(data)
    assert read.complete is True
    assert read.row_count == 2
    assert [r["Name"] for r in read.rows] == ["A", "B"]


def test_header_row_detected_after_title() -> None:
    """A single-cell title row before the header must be skipped."""
    data = _xlsx_bytes(
        [
            ["Monthly Export"],
            ["Name", "Price"],
            ["A", 1],
        ]
    )
    read = read_xlsx(data)
    assert read.complete is True
    assert read.headers == ["Name", "Price"]
    assert read.row_count == 1
    assert read.rows[0]["Name"] == "A"


def test_persian_headers_roundtrip() -> None:
    data = _xlsx_bytes([["نام", "قیمت"], ["آب معدنی", 500]])
    read = read_xlsx(data)
    assert read.complete is True
    assert read.headers == ["نام", "قیمت"]
    assert read.rows[0]["نام"] == "آب معدنی"


def test_unreadable_bytes_report_incomplete() -> None:
    read = read_xlsx(b"this is not an xlsx file at all")
    assert read.complete is False
    assert read.error is not None
    assert read.rows == []


def test_missing_sheet_reported_incomplete() -> None:
    data = _xlsx_bytes([["Name", "Price"], ["A", 1]], sheet="Main")
    read = read_xlsx(data, sheet_name="DoesNotExist")
    assert read.complete is False
    assert "DoesNotExist" in (read.error or "")


def test_sheet_selection() -> None:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "First"
    ws1.append(["A", "B"])
    ws1.append(["1", "2"])
    ws2 = wb.create_sheet("Second")
    ws2.append(["X", "Y"])
    ws2.append(["3", "4"])
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    read = read_xlsx(data, sheet_name="Second")
    assert read.complete is True
    assert read.headers == ["X", "Y"]
    assert read.rows[0]["X"] == "3"


def test_no_header_reports_incomplete() -> None:
    data = _xlsx_bytes([["only-one"], ["x"]])
    read = read_xlsx(data)
    assert read.complete is False
    assert read.error is not None


class _FakeSource:
    kind = "EXCEL_UPLOAD"
    sheet_name = None


def test_read_source_dispatch_excel() -> None:
    data = _xlsx_bytes([["Name", "Price"], ["A", 1]])
    read = read_source(_FakeSource(), file_bytes=data)
    assert read.complete is True
    assert read.rows[0]["Name"] == "A"


def test_read_source_excel_requires_file() -> None:
    read = read_source(_FakeSource(), file_bytes=None)
    assert read.complete is False
    assert read.error is not None
